"""Excel-versus-Python reconciliation.

The workbook is re-saved through openpyxl (which discards every cached formula
result), recalculated headlessly by LibreOffice, and read back. Every engine
line that has a worksheet row is then compared cell by cell, and the Checks
sheet is inspected. Requires ``openpyxl`` and a LibreOffice ``soffice`` binary.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .engine import run_model
from .excel import BuildInfo, build_workbook
from .inputs import Inputs
from .layout import FIRST_COL, SHEET_NAMES

_RECALC_ALWAYS = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" \
xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">\
<value>0</value></prop></item>
</oor:items>
"""


@dataclass
class ReconciliationReport:
    """Outcome of comparing one recalculated workbook scenario to the engine."""

    scenario: str
    compared: int = 0
    max_abs_diff: float = 0.0
    mismatches: list[str] = field(default_factory=list)
    check_cells: int = 0
    checks_pass: bool = False

    @property
    def ok(self) -> bool:
        return self.compared > 0 and not self.mismatches and self.checks_pass


def find_soffice() -> str | None:
    """Path to LibreOffice's ``soffice`` binary, or None when unavailable."""
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalculate(src: Path, outdir: Path, soffice: str | None = None, timeout: int = 240) -> Path:
    """Recalculate ``src`` with headless LibreOffice and return the converted xlsx path."""
    soffice = soffice or find_soffice()
    if soffice is None:
        raise RuntimeError("LibreOffice (soffice) is not installed")
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo_profile"
    (profile / "user").mkdir(parents=True, exist_ok=True)
    (profile / "user" / "registrymodifications.xcu").write_text(_RECALC_ALWAYS, encoding="utf-8")
    result_dir = outdir / "recalculated"
    result_dir.mkdir(exist_ok=True)
    cmd = [
        soffice,
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--headless",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(result_dir),
        str(src),
    ]
    subprocess.run(cmd, check=True, capture_output=True, timeout=timeout)
    out = result_dir / src.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice produced no output for {src}")
    return out


def _split_address(address: str) -> tuple[str, str]:
    sheet, cell = address.split("!")
    return sheet.strip("'"), cell.replace("$", "")


def reconcile(
    info: BuildInfo,
    inputs: Inputs,
    scenario: str,
    workdir: Path,
    edits: Mapping[tuple[str, str], Any] | None = None,
    soffice: str | None = None,
) -> ReconciliationReport:
    """Select ``scenario`` in the workbook, recalculate it, and compare with the engine.

    ``edits`` maps ``(sheet name, cell)`` to replacement input values; pass the
    same changes in ``inputs`` so the engine sees the same what-if.
    """
    import openpyxl  # imported lazily: only needed for validation

    workdir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(info.path)
    sheet, cell = _split_address(info.names["Scenario"])
    wb[sheet][cell] = scenario
    for (sheet_name, address), value in (edits or {}).items():
        wb[sheet_name][address] = value
    stripped = workdir / f"{scenario.lower()}_stripped.xlsx"
    wb.save(stripped)  # openpyxl keeps formulas but drops cached values
    recalculated = recalculate(stripped, workdir / scenario.lower(), soffice)
    values = openpyxl.load_workbook(recalculated, data_only=True)

    engine = run_model(inputs, scenario, "opening")
    report = ReconciliationReport(scenario=scenario)
    n_cols = len(engine.years)
    for _, sheet_name, row, key in info.recon:
        ws = values[sheet_name]
        for i in range(n_cols):
            got = ws.cell(row=row, column=FIRST_COL + 1 + i).value
            if isinstance(got, str) and got.startswith(("#", "Err:")):
                report.mismatches.append(f"{sheet_name} r{row} {key} col{i}: Excel error {got}")
                continue
            if got is None or isinstance(got, str):
                continue
            want = engine[key][i]
            report.compared += 1
            if want is None:
                report.mismatches.append(f"{sheet_name} r{row} {key} col{i}: Excel {got} but engine has no value")
                continue
            diff = abs(float(got) - want)
            report.max_abs_diff = max(report.max_abs_diff, diff)
            if diff > 1e-6 * max(1.0, abs(want)):
                report.mismatches.append(f"{sheet_name} r{row} {key} col{i}: Excel {got} vs engine {want}")

    checks = values[SHEET_NAMES["checks"]]
    flags: list[Any] = []
    for key, row in info.rows["checks"].items():
        if not key.startswith("chk_"):
            continue
        for col in range(FIRST_COL + 1, FIRST_COL + 1 + n_cols):
            v = checks.cell(row=row, column=col).value
            if isinstance(v, (int, float)):
                report.check_cells += 1
                if abs(v) > inputs.tolerance:
                    report.mismatches.append(f"Checks r{row} col{col}: non-zero check value {v}")
        flags.append(checks.cell(row=row, column=FIRST_COL + 1 + n_cols).value)
    sheet, cell = _split_address(info.names["Model_Checks_OK"])
    overall = values[sheet][cell].value
    report.checks_pass = bool(flags) and all(f is True or f == 1 for f in flags) and overall in (True, 1)
    return report


def validate_all(inputs: Inputs, scenarios: tuple[str, ...], workdir: Path | None = None) -> list[ReconciliationReport]:
    """Build the workbook once and reconcile each scenario through the dropdown."""
    owns_dir = workdir is None
    base = Path(tempfile.mkdtemp(prefix="three_statement_")) if workdir is None else workdir
    try:
        info = build_workbook(base / "model.xlsx", inputs, scenarios[0])
        return [reconcile(info, inputs, s, base / "work") for s in scenarios]
    finally:
        if owns_dir:
            shutil.rmtree(base, ignore_errors=True)
