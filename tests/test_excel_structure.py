"""Structure, formatting and formula-integrity tests (no LibreOffice needed)."""

from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from three_statement import run_model
from three_statement.excel import build_workbook
from three_statement.layout import SCALAR_NAMES, SHEET_NAMES

BLUE, BLACK, GREEN = "FF0000FF", "FF000000", "FF008000"
MODEL_SHEETS = ("Summary", "Income Statement", "Balance Sheet", "Cash Flow", "Schedules", "Checks")
FORECAST_COLS = range(6, 11)  # F..J
HIST_COLS = range(3, 6)  # C..E


@pytest.fixture(scope="module")
def built(tmp_path_factory, inputs):
    path = tmp_path_factory.mktemp("xlsx") / "model.xlsx"
    return build_workbook(path, inputs, "Base")


@pytest.fixture(scope="module")
def wb(built):
    return openpyxl.load_workbook(built.path)


@pytest.fixture(scope="module")
def wb_values(built):
    return openpyxl.load_workbook(built.path, data_only=True)


def _cells(ws, cols, first_row=6):
    for row in ws.iter_rows(min_row=first_row, min_col=cols[0], max_col=cols[-1]):
        for cell in row:
            if cell.value is not None:
                yield cell


def test_sheet_order_and_names(wb):
    assert wb.sheetnames == list(SHEET_NAMES.values())
    assert wb.sheetnames[0] == "Cover"


def test_named_ranges_exist_and_point_at_assumptions(wb, built):
    names = set(wb.defined_names)
    assert {"Scenario", "Scenario_Idx", "Days_In_Year", "Min_Cash", "Revolver_Limit", "Tolerance"} <= names
    assert "Model_Checks_OK" in names
    assert wb.defined_names["Scenario"].attr_text == "Assumptions!$C$7"
    assert wb.defined_names["Min_Cash"].attr_text.startswith("Assumptions!$C$")
    assert wb.defined_names["Model_Checks_OK"].attr_text.startswith("Checks!")


def test_scenario_dropdown_is_a_list_validation(wb):
    ws = wb["Assumptions"]
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    dv = validations[0]
    assert dv.type == "list" and "C7" in str(dv.sqref)
    assert dv.formula1 == "$C$9:$C$11"
    assert [ws[f"C{r}"].value for r in (9, 10, 11)] == ["Base", "Upside", "Downside"]


def test_scenario_index_uses_match_and_live_rows_use_choose(wb):
    ws = wb["Assumptions"]
    assert ws["C8"].value == "=MATCH(C7,C9:C11,0)"
    live = [c.value for c in ws["F"] if isinstance(c.value, str) and c.value.startswith("=CHOOSE")]
    assert len(live) == 8 and all("Scenario_Idx" in f for f in live)


def test_forecast_columns_contain_only_formulas_on_model_sheets(wb):
    for name in MODEL_SHEETS:
        for cell in _cells(wb[name], list(FORECAST_COLS)):
            assert not isinstance(cell.value, (int, float)), f"{name}!{cell.coordinate} is hard-coded"


def test_model_sheets_are_formula_driven(wb):
    counts = {
        n: sum(1 for c in _cells(wb[n], list(range(3, 12))) if str(c.value).startswith("=")) for n in MODEL_SHEETS
    }
    assert counts["Income Statement"] > 150 and counts["Balance Sheet"] > 110 and counts["Schedules"] > 200
    assert sum(counts.values()) > 900


def test_statements_pull_drivers_from_assumptions(wb, built):
    r = built.rows["is"]
    assert "Assumptions!" in wb["Income Statement"][f"F{r['rev_s1']}"].value
    assert "Assumptions!" in wb["Income Statement"][f"F{r['cogs']}"].value
    assert wb["Balance Sheet"][f"F{built.rows['bs']['cash']}"].value == f"='Cash Flow'!F{built.rows['cf']['cash_end']}"


def test_hard_coded_numbers_are_blue(wb):
    for name in (*MODEL_SHEETS, "Assumptions"):
        for cell in _cells(wb[name], list(range(3, 12))):
            if isinstance(cell.value, (int, float)):
                assert cell.font.color.rgb in (BLUE, "FF9DC3FF"), f"{name}!{cell.coordinate}"


def test_formula_colours_black_same_sheet_green_cross_sheet(wb):
    for name in (*MODEL_SHEETS, "Assumptions"):
        for cell in _cells(wb[name], list(range(1, 12))):
            f = cell.value
            if not (isinstance(f, str) and f.startswith("=")):
                continue
            external = "!" in f or (name != "Assumptions" and any(n in f for n in (*SCALAR_NAMES, "Scenario")))
            if cell.row == 5 or cell.row == 3:
                continue  # header and scenario banner
            want = GREEN if external else BLACK
            assert cell.font.color.rgb == want, f"{name}!{cell.coordinate} {f}"


def test_frozen_panes_and_hidden_gridlines(wb):
    for name in (*MODEL_SHEETS, "Assumptions"):
        assert wb[name].freeze_panes == "C6", name
    for ws in wb.worksheets:
        assert ws.sheet_view.showGridLines is False, ws.title


def test_print_setup_is_landscape_fit_to_width(wb):
    for name in ("Assumptions", "Income Statement", "Balance Sheet", "Cash Flow", "Schedules", "Checks"):
        ws = wb[name]
        assert ws.page_setup.orientation == "landscape"
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert ws.page_setup.fitToWidth in (None, 1)
        assert ws.print_area
        assert ws.print_title_rows == "$4:$5"
    assert wb["Summary"].page_setup.orientation == "portrait"


def test_years_are_column_headers_linked_to_assumptions(wb):
    ws = wb["Income Statement"]
    assert [ws.cell(row=5, column=c).value for c in (3, 8)] == ["=Assumptions!C$5", "=Assumptions!H$5"]
    assert ws["C5"].number_format == '"FY"0"A"' and ws["F5"].number_format == '"FY"0"E"'
    assert wb["Assumptions"]["C5"].value == 2023 and wb["Assumptions"]["D5"].value == "=C5+1"


def test_summary_contains_required_ratios(wb):
    labels = {c.value for c in wb["Summary"]["A"] if c.value}
    for needed in (
        "% margin",
        "EBITDA",
        "Return on equity",
        "Net debt / EBITDA",
        "Free cash flow",
        "Interest coverage (EBIT / interest expense)",
    ):
        assert needed in labels
    assert any(str(v).startswith("Return on invested capital") for v in labels)


def test_three_native_charts_on_summary(built):
    with zipfile.ZipFile(built.path) as z:
        charts = [n for n in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n)]
        assert len(charts) == 3
        sheet_rels = [n for n in z.namelist() if n.startswith("xl/worksheets/_rels/")]
        assert len(sheet_rels) == 1  # only Summary carries a drawing
        xml = " ".join(z.read(n).decode() for n in charts)
    assert "stackedBar" not in xml and "c:barChart" in xml and "c:lineChart" in xml


def test_checks_sheet_has_eleven_status_formulas(wb):
    ws = wb["Checks"]
    flags = [ws.cell(row=r, column=11).value for r in range(7, 18)]
    assert all(isinstance(f, str) and f.startswith("=SUMPRODUCT(ABS(") for f in flags)
    assert ws["C19"].value == "=AND(K7:K17)"


def test_cached_values_match_engine(built, wb_values, inputs):
    res = run_model(inputs, "Base")
    for _, sheet_name, row, key in built.recon:
        ws = wb_values[sheet_name]
        for i in range(8):
            cached = ws.cell(row=row, column=3 + i).value
            if isinstance(cached, (int, float)) and res[key][i] is not None:
                assert cached == pytest.approx(res[key][i], abs=1e-9)


def test_dropdown_scenario_and_cached_values_follow_build_argument(tmp_path, inputs):
    info = build_workbook(tmp_path / "d.xlsx", inputs, "downside")
    values = openpyxl.load_workbook(info.path, data_only=True)
    assert values["Assumptions"]["C7"].value == "Downside"
    assert values["Assumptions"]["C8"].value == 3
    row = info.rows["bs"]["revolver"]
    assert values["Balance Sheet"][f"J{row}"].value == pytest.approx(run_model(inputs, "Downside").last("revolver"))


def test_sensitivity_sheet_is_labelled_python_generated(wb, inputs):
    ws = wb["Sensitivity"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "PYTHON-GENERATED" in text and "static values" in text
    assert not any(str(c.value).startswith("=") for row in ws.iter_rows() for c in row if c.value)
    res = run_model(inputs, "Base")
    numbers = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, float)]
    assert any(v == pytest.approx(res.last("ebitda")) for v in numbers)
    assert any(v == pytest.approx(res.last("fcf")) for v in numbers)


def test_cover_has_colour_code_and_fictional_notice(wb):
    text = " ".join(str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value)
    for needle in ("FICTIONAL COMPANY", "Blue:", "Black:", "Green:", "Colour code", "Sheet guide"):
        assert needle in text
    assert wb["Cover"]["B30"].font.color.rgb == BLUE
    assert wb["Cover"]["B32"].font.color.rgb == GREEN


def test_document_properties_have_author_only(wb):
    assert wb.properties.creator == "Adham Karim"
    blob = " ".join(str(v) for v in (wb.properties.title, wb.properties.description, wb.properties.subject))
    assert "@" not in blob


# ------------------------------------------------------------------ circularity


_REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")


def _col_index(letters: str) -> int:
    return openpyxl.utils.column_index_from_string(letters)


def find_cycle(formulas: dict[tuple[str, str], str], names: dict[str, tuple[str, str]]):
    """Return a dependency cycle among formula cells, or None."""
    graph: dict[tuple[str, str], list[tuple[str, str]]] = {}
    for (sheet, coord), f in formulas.items():
        deps = []
        text = re.sub(r'"[^"]*"', "", f)
        for name, target in names.items():
            if re.search(rf"\b{name}\b", text):
                deps.append(target)
        for m in _REF.finditer(text):
            sh = m.group(1) or m.group(2) or sheet
            c1, r1, c2, r2 = m.group(3), int(m.group(4)), m.group(5), m.group(6)
            if c2:
                for ci in range(_col_index(c1), _col_index(c2) + 1):
                    for ri in range(r1, int(r2) + 1):
                        deps.append((sh, f"{get_column_letter(ci)}{ri}"))
            else:
                deps.append((sh, f"{c1}{r1}"))
        graph[(sheet, coord)] = deps
    state: dict[tuple[str, str], int] = {}
    for start in graph:
        if state.get(start):
            continue
        stack = [(start, iter(graph.get(start, [])))]
        state[start] = 1
        path = [start]
        while stack:
            node, it = stack[-1]
            nxt = next(it, None)
            if nxt is None:
                state[node] = 2
                stack.pop()
                path.pop()
                continue
            if state.get(nxt) == 1:
                return [*path[path.index(nxt) :], nxt]
            if not state.get(nxt) and nxt in graph:
                state[nxt] = 1
                path.append(nxt)
                stack.append((nxt, iter(graph[nxt])))
    return None


def test_cycle_detector_finds_a_cycle():
    cyc = {("S", "A1"): "=B1+1", ("S", "B1"): "=C1", ("S", "C1"): "=A1*2"}
    assert find_cycle(cyc, {}) is not None
    assert find_cycle({("S", "A1"): "=B1", ("S", "B1"): "=1"}, {}) is None


def test_workbook_has_no_circular_references(wb):
    formulas = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas[(ws.title, c.coordinate)] = c.value
    names = {}
    for name, dn in wb.defined_names.items():
        sheet, cell = dn.attr_text.split("!")
        names[name] = (sheet.strip("'"), cell.replace("$", ""))
    assert len(formulas) > 1000
    assert find_cycle(formulas, names) is None
