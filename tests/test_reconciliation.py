"""Excel versus Python reconciliation via headless LibreOffice recalculation.

Skipped automatically when LibreOffice (``soffice``) is not installed.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from three_statement import SCENARIOS, run_model
from three_statement.cli import main
from three_statement.excel import build_workbook
from three_statement.validate import find_soffice, reconcile

pytestmark = [
    pytest.mark.libreoffice,
    pytest.mark.skipif(find_soffice() is None, reason="LibreOffice (soffice) is not installed"),
]

KEY_LINES = {
    "Income Statement": ("revenue", "ebitda", "net_income"),
    "Balance Sheet": ("cash", "total_assets", "total_le", "revolver", "term_loan"),
    "Cash Flow": ("fcf",),
}
SHEET_KEY = {"Income Statement": "is", "Balance Sheet": "bs", "Cash Flow": "cf"}


@pytest.fixture(scope="module")
def workdir(tmp_path_factory) -> Path:
    return tmp_path_factory.mktemp("recon")


@pytest.fixture(scope="module")
def info(workdir, inputs):
    return build_workbook(workdir / "model.xlsx", inputs, "Base")


@pytest.fixture(scope="module")
def reports(info, inputs, workdir):
    return {s: reconcile(info, inputs, s, workdir / "work") for s in SCENARIOS}


def _recalculated(workdir: Path, scenario: str):
    path = workdir / "work" / scenario.lower() / "recalculated" / f"{scenario.lower()}_stripped.xlsx"
    return openpyxl.load_workbook(path, data_only=True)


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_every_mapped_line_matches_engine(reports, scenario):
    report = reports[scenario]
    assert report.mismatches == []
    assert report.compared > 1000  # ~1,070 formula cells per scenario
    assert report.max_abs_diff < 1e-6


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_key_line_items_match_engine(workdir, reports, info, inputs, scenario):
    assert reports[scenario].ok
    values = _recalculated(workdir, scenario)
    res = run_model(inputs, scenario)
    checked = 0
    for sheet, keys in KEY_LINES.items():
        ws = values[sheet]
        for key in keys:
            row = info.rows[SHEET_KEY[sheet]][key]
            for i in range(8):
                got = ws.cell(row=row, column=3 + i).value
                if res[key][i] is None:
                    assert got is None
                    continue
                assert got == pytest.approx(res[key][i], abs=0.01)
                assert got == pytest.approx(res[key][i], rel=1e-6, abs=1e-9)
                checked += 1
    assert checked >= 9 * 5


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_all_check_cells_are_zero_or_true(workdir, reports, info, scenario):
    assert reports[scenario].checks_pass and reports[scenario].check_cells > 50
    ws = _recalculated(workdir, scenario)["Checks"]
    for key, row in info.rows["checks"].items():
        if not key.startswith("chk_"):
            continue
        for col in range(3, 11):
            v = ws.cell(row=row, column=col).value
            assert v is None or abs(v) < 0.001
        assert ws.cell(row=row, column=11).value in (True, 1)
    assert ws["C19"].value in (True, 1)


def test_dropdown_changes_workbook_results(workdir, info):
    row = info.rows["bs"]["revolver"]
    base = _recalculated(workdir, "Base")["Balance Sheet"][f"J{row}"].value
    down = _recalculated(workdir, "Downside")["Balance Sheet"][f"J{row}"].value
    assert down > base + 10


def test_recalculation_is_real_not_cached(workdir, reports):
    stripped = openpyxl.load_workbook(workdir / "work" / "base_stripped.xlsx", data_only=True)
    assert stripped["Income Statement"]["F7"].value is None  # openpyxl wrote no cached value
    assert _recalculated(workdir, "Base")["Income Statement"]["F7"].value is not None


def _what_if(info, buyback: float, min_cash: float):
    min_cash_row = int(info.names["Min_Cash"].split("$")[-1])
    bb_row = info.rows["assump"]["buybacks"]
    edits = {("Assumptions", f"{col}{bb_row}"): buyback for col in "FGHIJ"}
    edits[("Assumptions", f"C{min_cash_row}")] = min_cash
    return edits


def test_what_if_edits_to_blue_inputs_flow_through(workdir, info, inputs):
    stressed = replace(inputs, buybacks=(60.0,) * 5, min_cash=80.0)
    report = reconcile(info, stressed, "Upside", workdir / "whatif", edits=_what_if(info, 60.0, 80.0))
    assert report.ok, report.mismatches[:3]
    assert max(run_model(stressed, "Upside").forecast("revolver")) > 100  # the revolver really is in play


def test_checks_flag_a_revolver_that_hits_its_commitment(workdir, info, inputs):
    stressed = replace(inputs, buybacks=(90.0,) * 5, min_cash=80.0)
    report = reconcile(info, stressed, "Upside", workdir / "breach", edits=_what_if(info, 90.0, 80.0))
    assert report.compared > 1000
    assert not report.checks_pass  # minimum cash can no longer be defended
    assert all(m.startswith("Checks") for m in report.mismatches)  # Excel still equals the engine


def test_cli_validate_command(capsys):
    assert main(["validate", "--scenario", "base"]) == 0
    out = capsys.readouterr().out
    assert "Base" in out and "OK" in out
